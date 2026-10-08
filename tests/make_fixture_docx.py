#!/usr/bin/env python3
"""Build a small, fictional Quicker template for the smoke test (no client data).
Usage: python3 make_fixture_docx.py <out.docx>   (needs python-docx)"""
import sys
from docx import Document

doc = Document()
for text in [
    "מספר שומה: {p.number}",
    "גוש {p.gush} חלקה {p.helka}",
    "תאריך הביקור: {p.ad.visitDate | date}",
    # anchor block: opener at the end of the paragraph before the block, closer at the end of its last paragraph
    "סוג הבנייה: {p.ad.buildingType}{#p.ad.buildingType == 'בנייה רוויה'}",
    "הנכס נמצא בבניין מגורים בבנייה רוויה.{/p.ad.buildingType == 'בנייה רוויה'}",
    "{#p.ad.hasElevator}בבניין יש מעלית.{/p.ad.hasElevator}{#!p.ad.hasElevator}בבניין אין מעלית.{/!p.ad.hasElevator}",
    "שווי השוק: {p.ad.marketValue | currency} ₪",
    "לווים:",
]:
    doc.add_paragraph(text)
t = doc.add_table(rows=2, cols=2)
t.rows[0].cells[0].text, t.rows[0].cells[1].text = "שם", "ת.ז"
t.rows[1].cells[0].text = "{#p.ad.borrowers}{borrowerName}"
t.rows[1].cells[1].text = "{borrowerId}{/p.ad.borrowers}"
doc.add_paragraph("{p.ad.photos | maxSize:300:200}")
doc.add_paragraph("הופק בתאריך {today}")
doc.save(sys.argv[1])
